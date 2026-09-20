#!/usr/bin/env python3
"""Read source cells as data. Never infer authorization, regions or date semantics."""
import argparse
import datetime
import hashlib
import json
from pathlib import Path
import openpyxl


def parse(path):
    workbook = openpyxl.load_workbook(path, data_only=True)
    sheet = workbook['학교입력 ']
    merged = {}
    for area in sheet.merged_cells.ranges:
        if area.min_col == area.max_col == 10:
            for row in range(area.min_row, area.max_row + 1):
                merged[row] = sheet.cell(area.min_row, 10).value
    def text(value):
        if value is None:
            return None
        if isinstance(value, (datetime.datetime, datetime.date)):
            return value.strftime('%Y-%m-%d')
        return str(value).replace('\r\n', '\n').replace('\r', '\n').strip()
    schools = []
    current = None
    for row in range(1, sheet.max_row + 1):
        values = [text(sheet.cell(row, col).value) for col in range(2, 11)]
        name, curriculum, region, authorized, target, exam, eligibility, notes, schedule = values
        if name == '학교명' or not any(values):
            current = None
            continue
        if name:
            if authorized not in (None, '인가', '비인가'):
                raise ValueError(f'Unsupported authorization at E{row}')
            current = dict(name=name, curriculumDescription=curriculum, region=region,
                           isAuthorized=None if authorized is None else authorized == '인가',
                           eligibility=eligibility, notes=notes, admissions=[])
            schools.append(current)
        if current is None:
            raise ValueError(f'Orphan row {row}')
        schedule = text(merged.get(row, sheet.cell(row, 10).value))
        if any((target, exam, schedule)):
            current['admissions'].append(dict(sourceRow=row, targetLabel=target, examContent=exam, scheduleText=schedule))
    if len(schools) != 18 or sum(len(s['admissions']) for s in schools) != 33:
        raise ValueError('Expected approved source: 18 schools / 33 admission rows')
    return dict(fileHash=hashlib.sha256(Path(path).read_bytes()).hexdigest(), sheet=sheet.title, schools=schools)


if __name__ == '__main__':
    parser = argparse.ArgumentParser()
    parser.add_argument('source')
    parser.add_argument('output')
    args = parser.parse_args()
    Path(args.output).write_text(json.dumps(parse(args.source), ensure_ascii=False, indent=2))
    Path(args.output).chmod(0o600)
    print('Validated 18 schools / 33 admission rows')
