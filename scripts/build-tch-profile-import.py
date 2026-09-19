#!/usr/bin/env python3
"""Build a teacher-only payload. Source workbook and generated JSON stay outside Git."""
import argparse
import hashlib
import json
import os
import re
from pathlib import Path
import openpyxl


def build(path):
    book = openpyxl.load_workbook(path, data_only=True)
    sheet = book['강사 정보']
    profiles = {
        str(row[0]).strip(): str(row[1]).strip().replace('\u2028', '\n')
        for row in book['현 강사 프로필'].iter_rows(min_row=2, max_row=26, values_only=True)
        if row[0] and row[1]
    }
    result = []
    clean = lambda value: str(value).strip() if value is not None else None
    for row in sheet.iter_rows(min_row=3, max_row=23, values_only=True):
        name = clean(row[1])
        if not name:
            continue
        contact = re.sub(r'[\u202a-\u202e\u2066-\u2069\n\r]', '', clean(row[7]) or '').strip()
        is_chat = contact.startswith('카톡ID:')
        subjects = clean(row[3])
        profile = profiles.get(name)
        tags = []
        for token, code in [('MAP', 'MAP'), ('SSAT', 'SSAT'), ('ISEE', 'ISEE')]:
            if re.search(r'\b' + token + r'\b', profile or '', re.I):
                tags.append(code)
        if '수학' in (subjects or ''):
            tags.append('MATH')
        held = []
        education = clean(row[2])
        if name == '조혜수':
            education = None
            held.append('education: conflicting source schools')
        if name == '최한나':
            subjects = None
            held.append('teachingSubjectsText: non-subject source value')
        result.append({'name': name, 'education': education,
            'teachingSubjectsText': subjects, 'experience': clean(row[6]),
            'residence': clean(row[5]), 'gender': {'남': 'MALE', '여': 'FEMALE'}.get(clean(row[4])),
            'phone': None if is_chat else contact or None,
            'kakaoId': contact.split(':', 1)[1].strip() if is_chat else None,
            'profileText': profile, 'subjects': tags, 'held': held})
    if len(result) != 21 or len({r['name'] for r in result}) != 21 or len(profiles) != 12:
        raise ValueError('Unexpected source counts; review workbook before import')
    return {'sourceSha256': hashlib.sha256(Path(path).read_bytes()).hexdigest(), 'teachers': result}


if __name__ == '__main__':
    parser = argparse.ArgumentParser()
    parser.add_argument('workbook')
    parser.add_argument('output')
    args = parser.parse_args()
    payload = build(args.workbook)
    fd = os.open(args.output, os.O_CREAT | os.O_EXCL | os.O_WRONLY, 0o600)
    with os.fdopen(fd, 'w') as output:
        json.dump(payload, output, ensure_ascii=False)
    print(json.dumps({'teachers': len(payload['teachers']), 'profiles': sum(bool(t['profileText']) for t in payload['teachers'])}))
